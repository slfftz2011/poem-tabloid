# -*- coding: utf-8 -*-
"""
build_ppt.py
============
诗歌小报「青春的另一种模样」生成脚本（python-pptx）

- 以 image1.png 为全屏背景（16:9，铺满不留边）
- 诗一《蕨生》逐行保留全角空格占位符，排成「青」字轮廓
- 诗二《静燃》逐行保留全角空格占位符，排成「春」字轮廓
- 两诗并排置于背景右上角月辉区（墨色文字，留白衬托）
- 严禁改写字句、严禁居中对齐，字形由每行内部与行前的全角空格（U+3000）构成

运行：python build_ppt.py
输出：poem-tabloid.pptx（与脚本同目录）
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
BG_IMG = os.path.join(BASE, "image1.png")
POEM1 = os.path.join(BASE, "poem1.md")
POEM2 = os.path.join(BASE, "poem2_revised.md")
OUT = os.path.join(BASE, "poem-tabloid.pptx")

INK = RGBColor(0x2B, 0x2B, 0x2B)      # 墨色
ACCENT = RGBColor(0x5A, 0x4A, 0x3A)   # 标题灰褐，与画面暖色协调


def load_poem_lines(path):
    """读取诗歌 md，去除 Markdown 标题行，将 &#12288; 实体还原为全角空格 U+3000。"""
    with open(path, encoding="utf-8") as f:
        raw = f.read().splitlines()
    lines = []
    for ln in raw:
        s = ln.strip()
        if not s or s.startswith("#"):
            continue
        s = s.replace("&#12288;", "\u3000")
        lines.append(s)
    return lines


def set_run_font(run, name="SimSun", size=13, color=INK):
    """设置 run 字体，同时覆盖西文与东亚字体（保证全角空格与汉字等宽）。"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def add_poem_shape(slide, x, y, w, h, title, lines, size=13):
    """添加一首诗的文字块：标题在上，字形正文逐行 LEFT 对齐，保留全角空格。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    # 标题行
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    p0.space_before = Pt(0)
    p0.space_after = Pt(6)
    p0.line_spacing = 1.0
    r0 = p0.add_run()
    r0.text = "《%s》" % title
    set_run_font(r0, size=size - 2, color=ACCENT)

    # 字形正文：每行一个段落，保持行内/行前全角空格
    for i, line in enumerate(lines):
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = line
        set_run_font(r, size=size)
    return tb


def main():
    poem1 = load_poem_lines(POEM1)
    poem2 = load_poem_lines(POEM2)
    print("蕨生 行数:", len(poem1), "最长:", max(len(l) for l in poem1))
    print("静燃 行数:", len(poem2), "最长:", max(len(l) for l in poem2))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式

    # 全屏背景：铺满不留边
    slide.shapes.add_picture(BG_IMG, Inches(0), Inches(0),
                             width=Inches(13.333), height=Inches(7.5))

    FONT = 13

    # 右上角月辉区：青（左）春（右）并排
    add_poem_shape(slide, x=8.65, y=1.35, w=1.75, h=4.2,
                   title="蕨生", lines=poem1, size=FONT)
    add_poem_shape(slide, x=10.40, y=1.35, w=2.60, h=4.2,
                   title="静燃", lines=poem2, size=FONT)

    prs.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
